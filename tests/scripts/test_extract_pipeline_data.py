"""Tests for ``scripts/extract_pipeline_data.py``.

The pipeline itself is not exercised here — we mock ``RAGPipeline``
to return a synthetic response/trace so the test runs offline (no
Ollama, no Qdrant).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

import pandas as pd
import pytest
from openpyxl import load_workbook

from scripts.extract_pipeline_data import (
    InputRow,
    explode_response,
    load_input,
    run_extraction,
    write_xlsx,
)


# ---------------------------------------------------------------------
# Synthetic pipeline response
# ---------------------------------------------------------------------


def _make_response(
    *,
    rewritten: List[Dict[str, Any]] = None,
    docs_per_query: Dict[str, List[Dict[str, Any]]] = None,
    evaluation_scores: List[Dict[str, Any]] = None,
    generator_documents: List[Dict[str, Any]] = None,
    answer: str = "Resposta final.",
    errors: List[str] = None,
) -> Dict[str, Any]:
    return {
        "answer": answer,
        "sources": [],
        "search_time_ms": 100,
        "llm_time_ms": 200,
        "total_time_ms": 300,
        "trace": {
            "original_query": "q",
            "rewritten_queries": rewritten or [],
            "search_documents_per_query": docs_per_query or {},
            "evaluation_scores": evaluation_scores or [],
            "generator_documents": generator_documents or [],
            "generator_model": "fake-model",
            "generator_grounded_only": True,
            "errors": errors or [],
            "timings": {
                "rewriter_ms": 10, "searcher_ms": 20,
                "evaluator_ms": 30, "generator_ms": 40, "total_ms": 100,
            },
        },
    }


# ---------------------------------------------------------------------
# load_input
# ---------------------------------------------------------------------


class TestLoadInput:
    def test_csv_basic(self, tmp_path: Path):
        path = tmp_path / "in.csv"
        path.write_text(
            "query_id,query,category\nQ1,oi,test\nQ2,tudo bem,prod\n",
            encoding="utf-8",
        )
        rows = load_input(path)
        assert [r.query_id for r in rows] == ["Q1", "Q2"]
        assert rows[0].extras == {"category": "test"}

    def test_xlsx_basic(self, tmp_path: Path):
        path = tmp_path / "in.xlsx"
        df = pd.DataFrame(
            [{"query_id": "Q1", "query": "oi", "expected": "DOC-1"}]
        )
        df.to_excel(path, index=False, engine="openpyxl")

        rows = load_input(path)
        assert len(rows) == 1
        assert rows[0].query_id == "Q1"
        assert rows[0].extras == {"expected": "DOC-1"}

    def test_missing_query_column_raises(self, tmp_path: Path):
        path = tmp_path / "bad.csv"
        path.write_text("question\noi\n", encoding="utf-8")
        with pytest.raises(ValueError, match="query"):
            load_input(path)

    def test_unsupported_format_raises(self, tmp_path: Path):
        path = tmp_path / "bad.txt"
        path.write_text("query\noi\n", encoding="utf-8")
        with pytest.raises(ValueError, match="Unsupported"):
            load_input(path)

    def test_auto_id_when_missing(self, tmp_path: Path):
        path = tmp_path / "noid.csv"
        path.write_text("query\nfoo\nbar\n", encoding="utf-8")
        rows = load_input(path)
        assert [r.query_id for r in rows] == ["Q001", "Q002"]

    def test_skips_blank_queries(self, tmp_path: Path):
        path = tmp_path / "blank.csv"
        path.write_text("query\nfoo\n\n   \nbar\n", encoding="utf-8")
        rows = load_input(path)
        assert [r.query for r in rows] == ["foo", "bar"]

    def test_zero_rows_raises(self, tmp_path: Path):
        path = tmp_path / "empty.csv"
        path.write_text("query\n   \n\n", encoding="utf-8")
        with pytest.raises(ValueError, match="No usable rows"):
            load_input(path)


# ---------------------------------------------------------------------
# explode_response
# ---------------------------------------------------------------------


class TestExplodeResponse:
    def _input_row(self, **extras: Any) -> InputRow:
        return InputRow(query_id="Q1", query="qual a regra?", extras=extras)

    def test_two_subqueries_three_docs_each_yields_six_rows(self):
        rewritten = [
            {"text": "sq1", "filters": [], "sorts": [], "facet_type": "general"},
            {"text": "sq2", "filters": [], "sorts": [], "facet_type": "temporal"},
        ]
        docs_per_query = {
            "sq1": [
                {"regulation_id": "d1", "url": "u1", "score": 0.9,
                 "type": "ICA", "number": "100-1", "title": "T1"},
                {"regulation_id": "d2", "url": "u2", "score": 0.8,
                 "type": "ICA", "number": "100-2", "title": "T2"},
                {"regulation_id": "d3", "url": "u3", "score": 0.7,
                 "type": "ICA", "number": "100-3", "title": "T3"},
            ],
            "sq2": [
                {"regulation_id": "d4", "url": "u4", "score": 0.95,
                 "type": "MCA", "number": "56-5", "title": "T4"},
                {"regulation_id": "d5", "url": "u5", "score": 0.85,
                 "type": "MCA", "number": "56-6", "title": "T5"},
                {"regulation_id": "d6", "url": "u6", "score": 0.75,
                 "type": "MCA", "number": "56-7", "title": "T6"},
            ],
        }
        evaluation_scores = [
            {"regulation_id": "d1", "score": 80.0, "accepted": True},
            {"regulation_id": "d4", "score": 90.0, "accepted": True},
            {"regulation_id": "d2", "score": 40.0, "accepted": False},
        ]
        generator_documents = [
            {"regulation_id": "d1", "type": "ICA", "number": "100-1",
             "title": "T1", "score": 80.0,
             "text": "Texto integral d1", "char_count": 17, "truncated": False},
            {"regulation_id": "d4", "type": "MCA", "number": "56-5",
             "title": "T4", "score": 90.0,
             "text": "Texto integral d4", "char_count": 17, "truncated": True},
        ]
        response = _make_response(
            rewritten=rewritten,
            docs_per_query=docs_per_query,
            evaluation_scores=evaluation_scores,
            generator_documents=generator_documents,
            answer="resposta XYZ",
        )

        rows = explode_response(self._input_row(category="test"), response)
        assert len(rows) == 6

        sent_rows = [r for r in rows if r["sent_to_generator"]]
        assert {r["regulation_id"] for r in sent_rows} == {"d1", "d4"}
        for r in sent_rows:
            assert r["generator_text"].startswith("Texto integral")

        not_sent = [r for r in rows if not r["sent_to_generator"]]
        for r in not_sent:
            assert r["generator_text"] == ""

        d1_row = next(r for r in rows if r["regulation_id"] == "d1")
        assert d1_row["evaluator_score"] == 80.0
        assert d1_row["evaluator_accepted"] is True
        assert d1_row["search_score"] == 0.9
        assert d1_row["subquery_text"] == "sq1"
        assert d1_row["doc_rank_in_subquery"] == 0
        assert d1_row["facet_type"] == "general"

        for r in rows:
            assert r["final_answer"] == "resposta XYZ"
            assert r["input__category"] == "test"
            assert r["query_id"] == "Q1"

    def test_zero_subqueries_yields_one_placeholder_row(self):
        response = _make_response(errors=["Rewriter: timeout"])
        rows = explode_response(self._input_row(), response)
        assert len(rows) == 1
        assert rows[0]["regulation_id"] == ""
        assert rows[0]["sent_to_generator"] is False
        assert "timeout" in rows[0]["errors"]

    def test_subquery_with_no_docs_keeps_placeholder_row(self):
        rewritten = [
            {"text": "sq1", "filters": [], "sorts": [], "facet_type": "general"},
        ]
        response = _make_response(
            rewritten=rewritten,
            docs_per_query={"sq1": []},
        )
        rows = explode_response(self._input_row(), response)
        assert len(rows) == 1
        assert rows[0]["subquery_text"] == "sq1"
        assert rows[0]["regulation_id"] == ""

    def test_filters_and_sorts_serialized_as_compact_json(self):
        rewritten = [{
            "text": "sq",
            "filters": [{"field": "metadata.type", "operator": "eq", "value": "ICA"}],
            "sorts": [{"field": "effective_date", "order": "desc"}],
            "facet_type": "document_type",
        }]
        docs_per_query = {"sq": [
            {"regulation_id": "d1", "url": "u1", "score": 0.5,
             "type": "ICA", "number": "100", "title": "T"},
        ]}
        response = _make_response(
            rewritten=rewritten, docs_per_query=docs_per_query,
        )
        rows = explode_response(self._input_row(), response)
        assert "metadata.type" in rows[0]["filters_json"]
        assert "effective_date" in rows[0]["sorts_json"]


# ---------------------------------------------------------------------
# write_xlsx
# ---------------------------------------------------------------------


class TestWriteXlsx:
    def test_writes_single_sheet_with_headers(self, tmp_path: Path):
        rows = [
            {"query_id": "Q1", "query": "oi", "regulation_id": "d1",
             "sent_to_generator": True, "generator_text": "abc"},
        ]
        out = tmp_path / "out.xlsx"
        write_xlsx(rows, out)

        assert out.exists()
        wb = load_workbook(out, read_only=True)
        assert wb.sheetnames == ["exploded"]
        sheet = wb["exploded"]
        header = [c.value for c in next(sheet.iter_rows(max_row=1))]
        assert header[0] == "query_id"
        assert "regulation_id" in header
        wb.close()

    def test_handles_empty_rows(self, tmp_path: Path):
        out = tmp_path / "empty.xlsx"
        write_xlsx([], out)
        assert out.exists()


# ---------------------------------------------------------------------
# run_extraction (end-to-end with mocked pipeline)
# ---------------------------------------------------------------------


class _FakePipeline:
    """Captures calls and returns canned responses."""

    def __init__(self, response_factory):
        self.calls: List[Dict[str, Any]] = []
        self._factory = response_factory

    def query(self, question: str, **kwargs: Any) -> Dict[str, Any]:
        self.calls.append({"question": question, **kwargs})
        return self._factory(question, kwargs)


class TestRunExtraction:
    def test_end_to_end_writes_expected_rows(self, tmp_path: Path):
        in_path = tmp_path / "in.csv"
        in_path.write_text(
            "query_id,query,category\nQ1,oi,test\nQ2,tchau,test\n",
            encoding="utf-8",
        )
        out_path = tmp_path / "out.xlsx"

        def factory(question: str, kwargs: Dict) -> Dict[str, Any]:
            return _make_response(
                rewritten=[{"text": question, "filters": [], "sorts": [],
                            "facet_type": "general"}],
                docs_per_query={question: [
                    {"regulation_id": f"d-{question}", "url": "u",
                     "score": 0.5, "type": "ICA", "number": "1",
                     "title": "T"},
                ]},
                generator_documents=[
                    {"regulation_id": f"d-{question}", "text": "TXT",
                     "char_count": 3, "truncated": False,
                     "type": "ICA", "number": "1", "title": "T",
                     "score": 80.0},
                ],
                answer=f"resp:{question}",
            )

        fake = _FakePipeline(factory)
        run_extraction(
            input_path=in_path, output_path=out_path,
            k=3, include_generation=True,
            pipeline_factory=lambda: fake,
        )

        assert len(fake.calls) == 2
        for call in fake.calls:
            assert call["debug"] is True
            assert call["limit"] == 3
            assert call["include_generation"] is True

        df = pd.read_excel(out_path, engine="openpyxl")
        assert len(df) == 2
        assert set(df["query_id"]) == {"Q1", "Q2"}
        assert all(df["sent_to_generator"])
        assert all(df["generator_text"] == "TXT")
        assert all(df["input__category"] == "test")

    def test_pipeline_crash_still_yields_row_with_error(self, tmp_path: Path):
        in_path = tmp_path / "in.csv"
        in_path.write_text("query\nfoo\n", encoding="utf-8")
        out_path = tmp_path / "out.xlsx"

        class _CrashPipeline:
            def query(self, question, **kwargs):
                raise RuntimeError("boom")

        run_extraction(
            input_path=in_path, output_path=out_path,
            k=1, pipeline_factory=lambda: _CrashPipeline(),
        )

        df = pd.read_excel(out_path, engine="openpyxl")
        assert len(df) == 1
        assert "boom" in df.iloc[0]["errors"]

    def test_sample_limits_rows(self, tmp_path: Path):
        in_path = tmp_path / "in.csv"
        in_path.write_text(
            "query\nq1\nq2\nq3\nq4\n", encoding="utf-8",
        )
        out_path = tmp_path / "out.xlsx"

        fake = _FakePipeline(
            lambda q, kw: _make_response(rewritten=[], errors=["x"]),
        )
        run_extraction(
            input_path=in_path, output_path=out_path,
            k=1, sample=2, pipeline_factory=lambda: fake,
        )
        assert len(fake.calls) == 2


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
